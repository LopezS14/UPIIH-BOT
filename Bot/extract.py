import fitz  # PyMuPDF
import streamlit as st
from docx import Document
import io

# Función para extraer texto del PDF
def extract_text_from_pdf(pdf_file):
    doc = fitz.open(stream=pdf_file.read(), filetype="pdf")
    text = ""
    for page_num in range(len(doc)):
        page = doc.load_page(page_num)
        text += page.get_text()
    return text

# Función para procesar el texto y extraer la información deseada
def extract_information(text):
    # Aquí puedes ajustar la lógica según el formato de tu PDF
    # Este es solo un ejemplo básico y necesitará ajustes para tu caso particular
    lines = text.split("\n")
    
    unidad = ""
    tema = ""
    temarios = []
    
    for line in lines:
        if "Unidad" in line:
            unidad = line.strip()
        elif "Tema" in line:
            tema = line.strip()
        elif "Temario" in line or line.startswith("-"):
            temarios.append(line.strip())
    
    return unidad, tema, temarios

# Función para modificar el archivo Word preestablecido
def fill_preformatted_word(template_path, unidad, tema, temarios):
    doc = Document(template_path)

    # Asumimos que el documento tiene placeholders como {Unidad}, {Tema} y {Temarios} donde queremos insertar la información.
    # Iteramos sobre los párrafos y reemplazamos los placeholders
    for paragraph in doc.paragraphs:
        if '{Unidad}' in paragraph.text:
            paragraph.text = paragraph.text.replace('{Unidad}', unidad)
        if '{Tema}' in paragraph.text:
            paragraph.text = paragraph.text.replace('{Tema}', tema)
        if '{Temarios}' in paragraph.text:
            # Añadimos los temarios en el mismo lugar
            temario_text = "\n".join(temarios)
            paragraph.text = paragraph.text.replace('{Temarios}', temario_text)

    # Guardar el documento en un objeto de memoria para permitir su descarga
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# Interfaz de Streamlit
st.title("Extractor de Información de PDF a Word con Formato Predefinido")

# Subir archivo PDF
pdf_file = st.file_uploader("Subir archivo PDF", type="pdf")

# Subir el archivo de Word preestablecido
word_template = st.file_uploader("Subir plantilla de Word", type="docx")

if pdf_file is not None and word_template is not None:
    # Extraer texto del PDF
    text = extract_text_from_pdf(pdf_file)
    
    # Procesar el texto para extraer la información deseada
    unidad, tema, temarios = extract_information(text)
    
    # Mostrar la información extraída
    st.subheader("Información extraída")
    st.write(f"Unidad: {unidad}")
    st.write(f"Tema: {tema}")
    st.write("Temarios:")
    for temario in temarios:
        st.write(f"- {temario}")
    
    # Generar archivo Word modificado
    if st.button("Generar archivo Word"):
        word_file = fill_preformatted_word(word_template, unidad, tema, temarios)
        st.download_button(
            label="Descargar archivo Word modificado",
            data=word_file,
            file_name="informacion_modificada.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
